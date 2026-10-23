"""Daily, weekly and monthly sales performance reports.

Laid out like the reports the team already circulates (Denri_Day_…,
Denri_Week_…, Denri_Month_… .xlsx): Summary, Executive Summary, Customer
Metrics, Store Performance, Gender Performance, Traffic, Revenue Analysis,
Data Quality, Feedback and Summary & Key Notes — each period set against the
one before it (the previous day, Sunday-to-Saturday week, or month).

Definitions, so every figure agrees with the rest of the dashboard:
  Revenue        — till sales as the Sales page counts them (every line but
                   delivery and customization fees, refunds netted).
  Corporate      — payments received on corporate invoices, as Odoo's Sales
                   & POS Dashboard shows them; Overall = Revenue + Corporate.
  Customers      — unique people, known by the last nine digits of their
                   phone; a shop's own number typed for a customer, or one
                   shared by many records, is not a customer.
  New / Repeat   — New: visited once in the period (one order); Repeat: two
                   or more orders in it — the Purchase Frequency table's
                   "1 Purchase" and "2+" rows. (The Shop Attendants score
                   cards count repeat purchase on distinct days instead.)
  Retention      — bought in this period and the one before, at the same shop.
  Traffic        — walk-ins counted (Footfall module, earlier gaps from the
                   saved Reporting-sheet history) against walk-in orders on the POS.
"""
from __future__ import annotations

import calendar
from dataclasses import dataclass
from datetime import date, datetime, timedelta

import numpy as np
import pandas as pd
import streamlit as st

from lib import feedback, names, phones, scorecards, staff_numbers

DAILY, WEEKLY, MONTHLY = "Daily", "Weekly", "Monthly"
CHANNELS = ["Walk-in", "Online", "Activation"]
FEEDBACK_LINK_TARGET = 75.0      # % of feedback links sent that come back answered
FLAG_CHANGE = 10.0               # % change worth flagging in the summary
SHOP_CHANGE = 40.0               # % change in a shop's revenue worth flagging
SHOP_MIN_REVENUE = 20_000        # ... when the shop sold at least this much
LOW_CONVERSION = 50.0            # walk-in conversion worth coaching below


# ---------------------------------------------------------------- periods ---

@dataclass
class Period:
    kind: str
    start: date
    end: date
    prev_start: date
    prev_end: date

    @property
    def label(self) -> str:
        if self.kind == DAILY:
            return f"{self.start:%A, %d %B %Y}"
        if self.kind == WEEKLY:
            return f"Week of {self.start:%d %b} - {self.end:%d %b %Y}"
        return f"{self.start:%B %Y}"

    @property
    def prev_label(self) -> str:
        if self.kind == DAILY:
            return f"{self.prev_start:%d %b %Y}"
        if self.kind == WEEKLY:
            return f"{self.prev_start:%d %b} - {self.prev_end:%d %b %Y}"
        return f"{self.prev_start:%d %b} - {self.prev_end:%d %b %Y}"

    @property
    def days(self) -> int:
        return (self.end - self.start).days + 1


def _month_end(d: date) -> date:
    return d.replace(day=calendar.monthrange(d.year, d.month)[1])


def period_for(kind: str, on: date) -> Period:
    if kind == DAILY:
        prev = on - timedelta(days=1)
        return Period(kind, on, on, prev, prev)
    if kind == WEEKLY:
        sunday = on - timedelta(days=(on.weekday() + 1) % 7)
        return Period(kind, sunday, sunday + timedelta(days=6),
                      sunday - timedelta(days=7), sunday - timedelta(days=1))
    first = on.replace(day=1)
    prev_end = first - timedelta(days=1)
    return Period(kind, first, _month_end(first), prev_end.replace(day=1), prev_end)


def bridge_periods(p: Period, count: int = 6) -> list[tuple[str, date, date]]:
    """The periods leading up to and including this one, oldest first: the
    last 7 days, 6 weeks or 6 months."""
    out = []
    if p.kind == DAILY:
        for i in range(6, -1, -1):
            d = p.start - timedelta(days=i)
            out.append((f"{d:%A, %d %B %Y}", d, d))
    elif p.kind == WEEKLY:
        for i in range(count - 1, -1, -1):
            s = p.start - timedelta(days=7 * i)
            e = s + timedelta(days=6)
            out.append((f"Week of {s:%d %b} - {e:%d %b %Y}", s, e))
    else:
        first = p.start
        months = []
        for _ in range(count):
            months.append(first)
            first = (first - timedelta(days=1)).replace(day=1)
        for m in reversed(months):
            out.append((f"{m:%B %Y}", m, _month_end(m)))
    return out


# ---------------------------------------------------------- formatting ---

def money(v) -> str:
    return "—" if v is None or pd.isna(v) else f"KES {v:,.2f}"


def change(cur, prev, *, money_fmt=False, pp=False, decimals=1) -> str:
    """'▲ +12.0%', '▼ -KES 1,200.00 (-3.0%)', '→ +0.0 pp' — as the reports write it."""
    if cur is None or pd.isna(cur):
        return "—"
    prev = 0 if prev is None or pd.isna(prev) else prev
    diff = cur - prev
    arrow = "▲" if diff > 1e-9 else "▼" if diff < -1e-9 else "→"
    if pp:
        return f"{arrow} {diff:+.{decimals}f} pp"
    pct = f"{diff / prev * 100:+.1f}%" if prev else "n/a"
    if money_fmt:
        sign = "+" if diff >= 0 else "-"
        return f"{arrow} {sign}KES {abs(diff):,.2f} ({pct})"
    return f"{arrow} {pct}"


def tone(cur, prev) -> str:
    if cur is None or prev is None or pd.isna(cur) or pd.isna(prev) or cur == prev:
        return "neutral"
    return "up" if cur > prev else "down"


def pct(part, whole) -> float:
    return float(part) / float(whole) * 100 if whole else 0.0


# ------------------------------------------------------------- customers ---

def prepare(orders: pd.DataFrame, not_customers: set[str]) -> pd.DataFrame:
    df = orders.copy()
    df["Revenue"] = pd.to_numeric(df["Revenue"], errors="coerce").fillna(0)
    df["Bags"] = pd.to_numeric(df["Bags"], errors="coerce").fillna(0)
    digits = df["Phone"].fillna("").astype(str).str.replace(r"\D", "", regex=True)
    df["Key"] = np.where(digits.str.len() >= 9, digits.str[-9:], None)
    # Not customers: store numbers, shared placeholders, and an attendant's
    # own number on any till but the Staff POS (where staff buy for themselves).
    attendant = (df["Key"].isin(staff_numbers.attendant_keys())
                 & (df["Location"] != staff_numbers.STAFF_POS))
    df.loc[df["Key"].isin(not_customers) | attendant, "Key"] = None
    df["Sale"] = df["Revenue"] > 0
    df["Gender"] = df["Gender"].replace({"Mwanamke": "Female", "Mwanaume": "Male"})
    return df


def placeholders(*frames: pd.DataFrame) -> set[str]:
    rows = []
    for f in frames:
        digits = f["Phone"].fillna("").astype(str).str.replace(r"\D", "", regex=True)
        rows.append(pd.DataFrame({"Customer": np.where(digits.str.len() >= 9, digits.str[-9:], None),
                                  "Partner": f["Partner"]}))
    return scorecards.placeholders(*[r.dropna(subset=["Customer"]) for r in rows], attendants=False)


def _mode(sales: pd.DataFrame, col: str) -> pd.Series:
    """Each customer's most frequent `col` (ties to the first alphabetically,
    as Series.mode does) — vectorised; a per-customer mode() took seconds."""
    counts = sales.groupby(["Key", col]).size().reset_index(name="n")
    return (counts.sort_values(["Key", "n", col], ascending=[True, False, True])
            .drop_duplicates("Key").set_index("Key")[col])


def customer_table(df: pd.DataFrame) -> pd.DataFrame:
    """One row per customer: orders, purchase days, revenue, shops, channel."""
    sales = df[df["Sale"] & df["Key"].notna()]
    if sales.empty:
        return pd.DataFrame(columns=["Key", "Orders", "Days", "Revenue", "Repeat"])
    g = sales.groupby("Key")
    # Revenue is net of the customer's refunds, so a mis-keyed order that was
    # reversed (e.g. 11,998,819 then -11,998,819) doesn't inflate their spend.
    net = df[df["Key"].isin(sales["Key"])].groupby("Key")["Revenue"].sum()
    out = pd.DataFrame({"Orders": g["Order ID"].nunique(), "Days": g["Date"].nunique(),
                        "Revenue": net,
                        "Gender": _mode(sales, "Gender"), "Location": _mode(sales, "Location")})
    out["Repeat"] = out["Orders"] >= 2          # visited more than once
    return out.reset_index()


def headline(df: pd.DataFrame, corporate: float) -> dict:
    cust = customer_table(df)
    sales = df[df["Sale"]]
    n = len(cust)
    by_channel = {c: sales[sales["Channel"] == c] for c in CHANNELS}
    records = df[df["Partner"].notna() & df["Sale"]]
    valid = records["Phone"].map(lambda p: phones.problem(p)[0] is None and bool(str(p).strip()))
    return {
        "Revenue": df["Revenue"].sum(),
        "Corporate": corporate,
        "Overall": df["Revenue"].sum() + corporate,
        "Orders": sales["Order ID"].nunique(),
        "Customers": n,
        "Repeat": int(cust["Repeat"].sum()) if n else 0,
        "New": n - (int(cust["Repeat"].sum()) if n else 0),
        "Repeat Rate": pct(cust["Repeat"].sum(), n) if n else 0.0,
        "Avg Rev/Txn": df["Revenue"].sum() / sales["Order ID"].nunique() if len(sales) else 0.0,
        "Avg Spend/Customer": cust["Revenue"].sum() / n if n else 0.0,
        "Walk-in Avg Spend": (df.loc[df["Channel"] == "Walk-in", "Revenue"].sum()
                              / max(by_channel["Walk-in"]["Order ID"].nunique(), 1)),
        "Online Avg Spend": (df.loc[df["Channel"] == "Online", "Revenue"].sum()
                             / max(by_channel["Online"]["Order ID"].nunique(), 1)),
        "Avg Orders": cust["Orders"].mean() if n else 0.0,
        "Data Quality": pct(valid.sum(), len(records)) if len(records) else 0.0,
        "Activation": pct(by_channel["Activation"]["Order ID"].nunique(), sales["Order ID"].nunique()),
    }


# --------------------------------------------------------------- sections ---

def executive(cur: dict, prev: dict, p: Period) -> tuple[pd.DataFrame, list[str], str]:
    rows = [
        ("Revenue", money(cur["Revenue"]), money(prev["Revenue"]), change(cur["Revenue"], prev["Revenue"])),
        ("Customers", f"{cur['Customers']:,}", f"{prev['Customers']:,}", change(cur["Customers"], prev["Customers"])),
        ("New", f"{cur['New']:,}", f"{prev['New']:,}", change(cur["New"], prev["New"])),
        ("Repeat", f"{cur['Repeat']:,}", f"{prev['Repeat']:,}", change(cur["Repeat"], prev["Repeat"])),
        ("Repeat Rate", f"{cur['Repeat Rate']:.1f}%", f"{prev['Repeat Rate']:.1f}%",
         change(cur["Repeat Rate"], prev["Repeat Rate"], pp=True)),
        ("Avg Rev/Txn", money(cur["Avg Rev/Txn"]), money(prev["Avg Rev/Txn"]),
         change(cur["Avg Rev/Txn"], prev["Avg Rev/Txn"])),
        ("Avg Spend/Customer", money(cur["Avg Spend/Customer"]), money(prev["Avg Spend/Customer"]),
         change(cur["Avg Spend/Customer"], prev["Avg Spend/Customer"])),
        ("Walk-in Avg Spend", money(cur["Walk-in Avg Spend"]), money(prev["Walk-in Avg Spend"]),
         change(cur["Walk-in Avg Spend"], prev["Walk-in Avg Spend"])),
        ("Online Avg Spend", money(cur["Online Avg Spend"]), money(prev["Online Avg Spend"]),
         change(cur["Online Avg Spend"], prev["Online Avg Spend"])),
        ("Avg Orders", f"{cur['Avg Orders']:.1f}", f"{prev['Avg Orders']:.1f}",
         change(cur["Avg Orders"], prev["Avg Orders"])),
        ("Data Quality", f"{cur['Data Quality']:.1f}%", f"{prev['Data Quality']:.1f}%",
         change(cur["Data Quality"], prev["Data Quality"], pp=True)),
        ("Activation", f"{cur['Activation']:.1f}%", f"{prev['Activation']:.1f}%",
         change(cur["Activation"], prev["Activation"], pp=True)),
    ]
    table = pd.DataFrame(rows, columns=["Metric", "Current", "Previous", "Change"])
    flags = []
    for key in ("Customers", "New", "Revenue", "Avg Spend/Customer", "Avg Rev/Txn", "Online Avg Spend"):
        c, pv = cur[key], prev[key]
        if pv and abs((c - pv) / pv * 100) >= FLAG_CHANGE:
            d = (c - pv) / pv * 100
            flags.append(f"[{'POSITIVE' if d > 0 else 'RISK'}] {key} is {'up' if d > 0 else 'down'} "
                         f"{abs(d):.1f}% vs {p.prev_label}.")
    d = (cur["Revenue"] - prev["Revenue"]) / prev["Revenue"] * 100 if prev["Revenue"] else 0
    word = "grew" if d > 2 else "softened" if d < -2 else "was broadly stable"
    note = (f"Revenue {word} {abs(d):.1f}% vs {p.prev_label}, with {cur['Customers']:,} customers "
            f"spending {money(cur['Avg Spend/Customer'])} on average.")
    return table, flags, note


def customer_metrics(cur_df: pd.DataFrame, prev_df: pd.DataFrame, cur: dict, prev: dict):
    metrics = pd.DataFrame([
        ("Total Customers", f"{cur['Customers']:,}", f"{prev['Customers']:,}", change(cur["Customers"], prev["Customers"])),
        ("New Customers", f"{cur['New']:,}", f"{prev['New']:,}", change(cur["New"], prev["New"])),
        ("Repeat Customers", f"{cur['Repeat']:,}", f"{prev['Repeat']:,}", change(cur["Repeat"], prev["Repeat"])),
        ("Repeat Customer Rate", f"{cur['Repeat Rate']:.1f}%", f"{prev['Repeat Rate']:.1f}%",
         change(cur["Repeat Rate"], prev["Repeat Rate"], pp=True)),
        ("Avg Orders / Customer", f"{cur['Avg Orders']:.2f}", f"{prev['Avg Orders']:.2f}",
         change(cur["Avg Orders"], prev["Avg Orders"])),
    ], columns=["Metric", "Current", "Previous", "Change"])

    cust = customer_table(cur_df)
    buckets = [("1 Purchase", 1, 1, "Primary re-engagement pool"),
               ("2 Purchases", 2, 2, "Early-loyalty tier"),
               ("3-5 Purchases", 3, 5, "Loyalty escalation target"),
               ("6+ Purchases", 6, 10 ** 9, "VIP cohort")]
    freq = pd.DataFrame([
        (label, int(((cust["Orders"] >= lo) & (cust["Orders"] <= hi)).sum()), read)
        for label, lo, hi, read in buckets], columns=["Purchase Frequency", "Customer Count", "Read"])
    freq.insert(2, "Percentage", (freq["Customer Count"] / max(len(cust), 1) * 100).round(1))

    sales = cur_df[cur_df["Sale"]]
    rows = []
    for c in CHANNELS:
        part = sales[sales["Channel"] == c]
        net = cur_df.loc[cur_df["Channel"] == c, "Revenue"].sum()
        rows.append((c, part["Key"].nunique(), net / max(part["Order ID"].nunique(), 1)))
    channel = pd.DataFrame(rows, columns=["Channel", "Customers", "Avg Spend"])
    channel["Share %"] = (channel["Customers"] / max(channel["Customers"].sum(), 1) * 100).round(1)
    channel["Avg Spend"] = channel["Avg Spend"].round(2)
    top = channel.loc[channel["Avg Spend"].idxmax(), "Channel"] if channel["Avg Spend"].max() else None
    channel["Note"] = np.where(channel["Channel"] == top, "Highest avg spend", "")
    channel = channel[["Channel", "Customers", "Share %", "Avg Spend", "Note"]]
    d = cur["Repeat Rate"] - prev["Repeat Rate"]
    note = (f"Repeat customer rate {'rose' if d > 0.5 else 'fell' if d < -0.5 else 'held steady'} "
            f"at {cur['Repeat Rate']:.1f}% ({d:+.1f} pp).")
    return metrics, freq, channel, note


def store_performance(cur_df: pd.DataFrame, prev_df: pd.DataFrame):
    cur_rev = cur_df.groupby("Location")["Revenue"].sum()
    prev_rev = prev_df.groupby("Location")["Revenue"].sum()
    known = cur_df[cur_df["Sale"] & cur_df["Key"].notna()]
    customers = known.groupby("Location")["Key"].nunique()
    # A repeat customer of a shop placed two or more orders there.
    rep = known.groupby(["Location", "Key"])["Order ID"].nunique().ge(2).groupby("Location").sum()
    shops = sorted(set(cur_rev.index) | set(prev_rev.index))
    table = pd.DataFrame({"Shop": shops})
    table["Current"] = table["Shop"].map(cur_rev).fillna(0).round(2)
    table["Previous"] = table["Shop"].map(prev_rev).fillna(0).round(2)
    table["Change"] = (table["Current"] - table["Previous"]).round(2)
    table["Change %"] = (table["Change"] / table["Previous"].where(table["Previous"] != 0) * 100).round(1)
    table["Customers"] = table["Shop"].map(customers).fillna(0).astype(int)
    table["Repeat"] = table["Shop"].map(rep).fillna(0).astype(int)
    table["New"] = table["Customers"] - table["Repeat"]
    table["Repeat Rate %"] = (table["Repeat"] / table["Customers"].where(table["Customers"] > 0)
                              * 100).round(1).fillna(0)
    table = table[["Shop", "Current", "Previous", "Change", "Change %", "New", "Repeat",
                   "Repeat Rate %"]].sort_values("Current", ascending=False)

    sales = cur_df[cur_df["Sale"]]
    mix = sales.pivot_table(index="Location", columns="Channel", values="Order ID",
                            aggfunc="nunique", fill_value=0)
    val = (cur_df.pivot_table(index="Location", columns="Channel", values="Revenue",
                              aggfunc="sum", fill_value=0)
           .reindex(mix.index, fill_value=0))
    out = pd.DataFrame(index=mix.index)
    for c in CHANNELS:
        out[c] = mix[c] if c in mix else 0
        out[f"{c} Sales (KES)"] = val[c].round(2) if c in val else 0.0
    out["Total"] = mix.sum(axis=1)
    out["Total Sales (KES)"] = val.sum(axis=1).round(2)
    out["Online %"] = (out["Online"] / out["Total"].where(out["Total"] > 0) * 100).round(1)
    out = out.reset_index().rename(columns={"Location": "Shop"}).sort_values("Total", ascending=False)
    return table, out


def retention_table(cur_df: pd.DataFrame, prev_df: pd.DataFrame) -> pd.DataFrame:
    """Per shop: this period's customers and sales, the previous period's,
    repeat customers (2+ orders this period) and retained customers (also bought
    at the shop last period), with their sales."""
    def by_shop(df):
        s = df[df["Sale"] & df["Key"].notna()]
        orders = s.groupby(["Location", "Key"])["Order ID"].nunique()
        net = df[df["Key"].notna()].groupby(["Location", "Key"])["Revenue"].sum()
        return pd.DataFrame({"Orders": orders, "Revenue": net.reindex(orders.index)})
    cur, prev = by_shop(cur_df), by_shop(prev_df)
    cur["Retained"] = cur.index.isin(prev.index)
    g = cur.groupby(level="Location")
    pg = prev.groupby(level="Location")
    shops = sorted(set(g.groups) | set(pg.groups))
    t = pd.DataFrame(index=shops)
    t["Customers"] = g.size()
    t["Sales (KES)"] = g["Revenue"].sum()
    t["Previous"] = pg.size()
    t["Previous Sales (KES)"] = pg["Revenue"].sum()
    t = t.fillna(0)
    t["Change"] = t["Customers"] - t["Previous"]
    rep = cur[cur["Orders"] >= 2].groupby(level="Location")
    t["Repeat"] = rep.size()
    t["Repeat Sales (KES)"] = rep["Revenue"].sum()
    ret = cur[cur["Retained"]].groupby(level="Location")
    t["Retention"] = ret.size()
    t["Retention Sales (KES)"] = ret["Revenue"].sum()
    t = t.fillna(0)
    t["Repeat Rate %"] = (t["Repeat"] / t["Customers"].where(t["Customers"] > 0) * 100).round(1).fillna(0)
    t["Retention Rate %"] = (t["Retention"] / t["Customers"].where(t["Customers"] > 0) * 100).round(1).fillna(0)
    for c in ("Customers", "Previous", "Change", "Repeat", "Retention"):
        t[c] = t[c].astype(int)
    for c in [c for c in t if c.endswith("(KES)")]:
        t[c] = t[c].round(2)
    t = t.reset_index().rename(columns={"index": "Shop"}).sort_values("Customers", ascending=False)
    return t[["Shop", "Customers", "Sales (KES)", "Previous", "Previous Sales (KES)", "Change",
              "Repeat", "Repeat Sales (KES)", "Repeat Rate %", "Retention",
              "Retention Sales (KES)", "Retention Rate %"]]


GENDERS = ["Female", "Male", "Corporate"]


def gender_split(cur_df: pd.DataFrame, prev_df: pd.DataFrame, p: Period):
    cur, prev = customer_table(cur_df), customer_table(prev_df)

    def share(c):
        counts = c["Gender"].value_counts() if len(c) else pd.Series(dtype=int)
        known = counts.reindex(GENDERS).fillna(0)
        return known, known / max(known.sum(), 1) * 100
    cn, cp = share(cur)
    pn, pp_ = share(prev)
    overall = pd.DataFrame({"Gender": GENDERS, "Count": cn.astype(int).values,
                            "%": cp.round(1).values,
                            "Change": [change(cp[g], pp_[g], pp=True) for g in GENDERS]})
    by_shop = (cur[cur["Gender"].isin(GENDERS)]
               .pivot_table(index="Location", columns="Gender", values="Key", aggfunc="count",
                            fill_value=0)
               .reindex(columns=GENDERS, fill_value=0))
    by_shop["Total"] = by_shop.sum(axis=1)
    by_shop = by_shop.reset_index().rename(columns={"Location": "Shop"}).sort_values("Total", ascending=False)
    ratio = f"{cn['Female'] / cn['Male']:.2f} : 1 (Female : Male)" if cn["Male"] else "—"
    shift = cp["Female"] - pp_["Female"]
    note = (f"Notable gender-mix shift: Female share {'rose' if shift > 0 else 'fell'} "
            f"{abs(shift):.1f} pp vs {p.prev_label}." if abs(shift) >= 5
            else f"Gender mix was stable vs {p.prev_label} - no shifts of note.")
    unknown = int((~cur["Gender"].isin(GENDERS)).sum()) if len(cur) else 0
    return ratio, overall, by_shop, note, unknown, shift


def traffic(foot: pd.DataFrame, cur_df: pd.DataFrame) -> tuple[pd.DataFrame, str]:
    """Walk-ins counted against walk-in purchases, and online/activation orders."""
    sales = cur_df[cur_df["Sale"]]
    orders = sales.pivot_table(index="Location", columns="Channel", values="Order ID",
                               aggfunc="nunique", fill_value=0)
    counted = foot[foot["Visitors"].notna()]
    t = pd.DataFrame({
        "Walk-in Purchased": counted.groupby("Location")["Walk-in Orders"].sum(),
        "Walk-in Total": counted.groupby("Location")["Visitors"].sum(),
    })
    t = t.join(orders.reindex(columns=["Online", "Activation"], fill_value=0), how="outer").fillna(0)
    t["Total Customers"] = sales[sales["Key"].notna()].groupby("Location")["Key"].nunique()
    t = t.fillna(0)
    t["Conv. Rate %"] = (t["Walk-in Purchased"] / t["Walk-in Total"].where(t["Walk-in Total"] > 0)
                         * 100).round(1)
    t = t[t["Walk-in Total"] > 0]
    for c in ("Walk-in Purchased", "Walk-in Total", "Online", "Activation", "Total Customers"):
        t[c] = t[c].astype(int)
    t = (t.reset_index().rename(columns={"index": "Shop", "Location": "Shop"})
         [["Shop", "Walk-in Purchased", "Walk-in Total", "Conv. Rate %", "Online", "Activation",
           "Total Customers"]].sort_values("Walk-in Purchased", ascending=False))
    note = ""
    if not t.empty:
        low = t.sort_values("Conv. Rate %").iloc[0]
        note = (f"{low['Shop']}'s conversion rate ({low['Conv. Rate %']:.1f}%) is the lowest - "
                f"{int(low['Walk-in Total']):,} walk-ins, {int(low['Walk-in Purchased']):,} bought.")
    return t, note


SHARED = "Shared"


def data_quality(cur_df: pd.DataFrame, prev_df: pd.DataFrame, not_customers: set[str]):
    """Phone-number quality, and the sales that carry a number that isn't the
    customer's: a store's own number, an attendant's own number (except at the
    Staff POS), or a number shared by many customer records. Store and
    attendant numbers come from Odoo — see lib.staff_numbers."""
    staff = staff_numbers.load().set_index("Key")
    stores, attendants = staff_numbers.store_keys(), staff_numbers.attendant_keys()

    def stats(df):
        rec = df[df["Partner"].notna() & df["Sale"]]
        issue = rec["Phone"].map(lambda p: phones.problem(p)[0] if str(p).strip() else "N/A")
        digits = rec["Phone"].fillna("").astype(str).str.replace(r"\D", "", regex=True).str[-9:]
        kind = pd.Series(None, index=rec.index, dtype=object)
        kind[digits.isin(not_customers)] = SHARED
        kind[digits.isin(attendants) & (rec["Location"] != staff_numbers.STAFF_POS)] = staff_numbers.ATTENDANT
        kind[digits.isin(stores)] = staff_numbers.STORE
        flagged = rec.assign(Digits=digits, Kind=kind)[kind.notna()]
        store_like = flagged["Kind"].isin([staff_numbers.STORE, SHARED])
        return {"records": len(rec), "na": int((issue == "N/A").sum()),
                "short": int(issue.fillna("").str.startswith(("Too short", "Missing")).sum()),
                "prefix": int(issue.fillna("").str.startswith(("Doesn't", "Starts with 00")).sum()),
                "long": int(issue.fillna("").str.startswith(("Too long", "International")).sum()),
                "valid": int(issue.isna().sum()),
                "store_unique": int(flagged.loc[store_like, "Digits"].nunique()),
                "attendant_unique": int(flagged.loc[~store_like, "Digits"].nunique()),
                "flagged": flagged}
    c, p = stats(cur_df), stats(prev_df)
    invalid = lambda s: s["short"] + s["prefix"] + s["long"]
    pc = lambda n, s: f"{n / s['records'] * 100:.1f}%" if s["records"] else "—"
    rows = [
        ("Total Records", f"{c['records']:,}", "-", f"{p['records']:,}", "-", change(c["records"], p["records"])),
        ("Valid Phone Numbers", f"{c['valid']:,}", pc(c["valid"], c), f"{p['valid']:,}", pc(p["valid"], p),
         change(c["valid"], p["valid"])),
        ("N/A Phones", f"{c['na']:,}", pc(c["na"], c), f"{p['na']:,}", pc(p["na"], p), change(c["na"], p["na"])),
        ("Invalid", f"{invalid(c):,}", pc(invalid(c), c), f"{invalid(p):,}", pc(invalid(p), p),
         change(invalid(c), invalid(p))),
        ("↳ Incomplete (under 10 digits)", f"{c['short']:,}", pc(c["short"], c), f"{p['short']:,}",
         pc(p["short"], p), change(c["short"], p["short"])),
        ("↳ Wrong prefix", f"{c['prefix']:,}", pc(c["prefix"], c), f"{p['prefix']:,}", pc(p["prefix"], p),
         change(c["prefix"], p["prefix"])),
        ("↳ Too long", f"{c['long']:,}", pc(c["long"], c), f"{p['long']:,}", pc(p["long"], p),
         change(c["long"], p["long"])),
        ("Store Number Used (unique)", f"{c['store_unique']:,}", pc(c["store_unique"], c),
         f"{p['store_unique']:,}", pc(p["store_unique"], p), change(c["store_unique"], p["store_unique"])),
        ("Attendant Number Used (unique)", f"{c['attendant_unique']:,}", pc(c["attendant_unique"], c),
         f"{p['attendant_unique']:,}", pc(p["attendant_unique"], p),
         change(c["attendant_unique"], p["attendant_unique"])),
    ]
    table = pd.DataFrame(rows, columns=["Metric", "Current", "Current %", "Previous", "Previous %", "Change"])

    flagged = c["flagged"].copy()
    flagged["Group"] = np.where(flagged["Kind"] == staff_numbers.ATTENDANT, "Attendant", "Store")
    by_shop = pd.DataFrame(index=pd.Index(sorted(flagged["Location"].unique()), name="Location"))
    for g in ("Store", "Attendant"):
        part = flagged[flagged["Group"] == g].groupby("Location")["Digits"]
        by_shop[f"{g} Numbers"] = part.nunique().reindex(by_shop.index).fillna(0).astype(int)
        by_shop[f"{g} Uses"] = part.size().reindex(by_shop.index).fillna(0).astype(int)
    by_shop["Total Uses"] = by_shop["Store Uses"] + by_shop["Attendant Uses"]
    by_shop = (by_shop.reset_index().rename(columns={"Location": "Shop"})
               .sort_values(["Total Uses", "Shop"], ascending=[False, True]))

    def belongs_to(kind, digits):
        if kind == SHARED or digits not in staff.index:
            return "Shared by many customer records"
        who = staff.loc[digits]
        if kind == staff_numbers.STORE:
            return f"{who['Owner']} store"
        return f"{who['Owner']} ({who['Shop']})"
    detail = (flagged.groupby(["Location", "Kind", "Digits"]).size().reset_index(name="Uses")
              .sort_values(["Uses", "Location"], ascending=[False, True]))
    detail["Belongs To"] = [belongs_to(k, d) for k, d in zip(detail["Kind"], detail["Digits"])]
    detail["Number"] = [staff.loc[d, "Number"] if d in staff.index else f"…{d}" for d in detail["Digits"]]
    detail = detail.rename(columns={"Location": "Shop"})[["Shop", "Kind", "Belongs To", "Number", "Uses"]]

    health = c["valid"] / c["records"] * 100 if c["records"] else 0
    note = (f"KYC health score: {health:.1f}% of records carry a valid phone number. Store and "
            "attendant numbers are flagged wherever they stand in for a customer's; an "
            "attendant's own number is only theirs to use at the Staff POS.")
    notes = []
    for g, text in (("Store", "Shop's own number used in place of a customer's at"),
                    ("Attendant", "Attendants' own numbers printed on customer receipts at")):
        uses = by_shop[by_shop[f"{g} Uses"] > 0]
        if not uses.empty:
            notes.append(f"{text}: " + ", ".join(f"{shop} ({n})" for shop, n
                                                in zip(uses["Shop"], uses[f"{g} Uses"])) + ".")
    return table, by_shop, detail, note, notes


def _branch(value: str) -> str:
    text = str(value or "").strip()
    low = text.lower()
    if low in ("ktda", "ktda shop"):
        return "Ktda"
    if "dar" in low or low == "sinza":
        return "Sinza"
    if "website" in low or "online" in low:
        return "Website"
    return text.title() if text else "Not recorded"


def feedback_section(responses: pd.DataFrame | None, links: pd.DataFrame | None, p: Period):
    """Ratings and responses by shop, this period against the last, and (when
    the WhatsApp Monitor's links are known) online responses against links sent."""
    if responses is None or responses.empty:
        return None
    r = responses.copy()
    r["Day"] = r["Submitted"].dt.date
    r["Shop"] = r["Branch"].map(_branch)
    for col in ("Experience", "Assistance"):
        r[col] = pd.to_numeric(r[col], errors="coerce")
    cur = r[(r["Day"] >= p.start) & (r["Day"] <= p.end)]
    prev = r[(r["Day"] >= p.prev_start) & (r["Day"] <= p.prev_end)]

    def status(v, n):
        return "Excellent" if v >= 4.5 else "Good" if v >= 4 else "Watch"
    rows = []
    # Rated from the main form only and counted over every form, as the
    # feedback sheet's own Report tab does.
    rated = lambda f: f[f["Form"] == feedback.RATED_FORM]
    for label, col in (("Overall Experience", "Experience"), ("Professionalism", "Assistance")):
        c, pv = rated(cur)[col].mean(), rated(prev)[col].mean()
        rows.append((label, round(c, 2) if pd.notna(c) else None, round(pv, 2) if pd.notna(pv) else None,
                     change(c, pv), status(c, len(cur)) if pd.notna(c) else "—"))
    d = len(cur) - len(prev)
    rows.append(("Feedback Responses", len(cur), len(prev), change(len(cur), len(prev)),
                 "Watch" if d < 0 else "Good"))
    summary = pd.DataFrame(rows, columns=["Metric", "Current", "Previous", "Change %", "Status"])
    shops = sorted(set(cur["Shop"]) | set(prev["Shop"]))
    by_shop = pd.DataFrame({"Shop": shops})
    by_shop["Current"] = by_shop["Shop"].map(cur["Shop"].value_counts()).fillna(0).astype(int)
    by_shop["Previous"] = by_shop["Shop"].map(prev["Shop"].value_counts()).fillna(0).astype(int)
    by_shop["Change"] = [change(a, b) for a, b in zip(by_shop["Current"], by_shop["Previous"])]
    by_shop = by_shop.sort_values("Current", ascending=False)

    targets = None
    if links is not None and not links.empty:
        src = cur["Source"].fillna("").str.lower()
        online = cur[src.str.contains("online")].groupby("Shop").size()
        walkin = cur[src.str.contains("walk")].groupby("Shop").size()
        sent = links.groupby("Branch").size()
        targets = pd.DataFrame({"Links Sent": sent})
        targets["Online"] = online.reindex(targets.index).fillna(0).astype(int)
        targets["Achv %"] = (targets["Online"] / targets["Links Sent"] * 100).round(1)
        targets["Gap (pp)"] = (targets["Achv %"] - FEEDBACK_LINK_TARGET).round(1)
        targets["Status"] = np.where(targets["Achv %"] >= FEEDBACK_LINK_TARGET, "On Target",
                                     np.where(targets["Achv %"] >= 50, "Below Target", "Critical"))
        targets["Walk-Ins"] = walkin.reindex(targets.index).fillna(0).astype(int)
        targets = (targets.reset_index().rename(columns={"index": "Shop", "Branch": "Shop"})
                   .sort_values("Links Sent", ascending=False))
    return {"summary": summary, "by_shop": by_shop, "targets": targets,
            "responses": len(cur), "previous": len(prev)}


def key_notes(p: Period, cur: dict, prev: dict, stores: pd.DataFrame, traffic_t: pd.DataFrame,
              gender_shift: float, fb: dict | None, data_notes: list[str]) -> dict[str, list[str]]:
    notes = {"Marketing": [], "Sales": [], "Data": []}
    d = (cur["Revenue"] - prev["Revenue"]) / prev["Revenue"] * 100 if prev["Revenue"] else 0
    if d <= -FLAG_CHANGE:
        notes["Marketing"].append(f"[CRITICAL] Revenue declined {abs(d):.1f}% vs {p.prev_label} - "
                                  "review demand drivers and promotions.")
    dn = (cur["New"] - prev["New"]) / prev["New"] * 100 if prev["New"] else 0
    if dn <= -FLAG_CHANGE:
        notes["Marketing"].append(f"[CRITICAL] New customer acquisition is down {abs(dn):.1f}% vs "
                                  f"{p.prev_label} - strengthen acquisition campaigns.")
    if abs(gender_shift) >= 5:
        notes["Marketing"].append(f"[MEDIUM] Female share {'rose' if gender_shift > 0 else 'fell'} "
                                  f"{abs(gender_shift):.1f} pp - check the campaign mix.")
    if fb and fb.get("targets") is not None and not fb["targets"].empty:
        t = fb["targets"]
        rate = t["Online"].sum() / max(t["Links Sent"].sum(), 1) * 100
        if rate < FEEDBACK_LINK_TARGET:
            notes["Marketing"].append(f"[CRITICAL] Online feedback response rate is {rate:.1f}% against "
                                      f"the {FEEDBACK_LINK_TARGET:.0f}% target.")
    big = stores[(stores["Current"] >= SHOP_MIN_REVENUE) | (stores["Previous"] >= SHOP_MIN_REVENUE)]
    for _, r in big.sort_values("Change %").iterrows():
        if pd.notna(r["Change %"]) and r["Change %"] <= -SHOP_CHANGE:
            notes["Sales"].append(f"[CRITICAL] {r['Shop']} revenue is down {abs(r['Change %']):.1f}% "
                                  f"vs {p.prev_label} - review staffing, stock and footfall.")
            break
    if not traffic_t.empty:
        low = traffic_t.sort_values("Conv. Rate %").iloc[0]
        if low["Conv. Rate %"] < LOW_CONVERSION:
            notes["Sales"].append(f"[HIGH] {low['Shop']} conversion rate is {low['Conv. Rate %']:.1f}% "
                                  "- coach staff on walk-in-to-sale conversion.")
    for _, r in big.sort_values("Change %", ascending=False).iterrows():
        if pd.notna(r["Change %"]) and r["Change %"] >= SHOP_CHANGE:
            notes["Sales"].append(f"[WIN] {r['Shop']} revenue is up {r['Change %']:.1f}% vs "
                                  f"{p.prev_label} - identify and repeat what worked.")
            break
    for text in data_notes:
        notes["Data"].append(f"[MEDIUM] {text}")
    return notes


# ----------------------------------------------------------------- build ---

CLOSED_TTL = 12 * 3600      # a finished day, week or month no longer changes
OPEN_TTL = 600              # the period still running refreshes every 10 minutes


@st.cache_data(ttl=CLOSED_TTL, show_spinner=False, max_entries=400)
def _orders_closed(start: date, end: date) -> pd.DataFrame:
    from lib import db, queries
    return db.run_query(queries.REPORT_ORDERS, {"start_date": start, "end_date": end})


def _orders(start: date, end: date) -> pd.DataFrame:
    """A period's orders. A period that has ended is kept for CLOSED_TTL, so
    the six weeks behind next week's bridge, or last month's report opened by
    someone else, come from memory rather than the replica."""
    from lib import db, queries
    if end < date.today():
        return _orders_closed(start, end)
    return db.run_query(queries.REPORT_ORDERS, {"start_date": start, "end_date": end})


def _corporate(start: date, end: date) -> float:
    from lib import db, queries
    t = db.run_query(queries.CORPORATE_TOTALS, {"start_date": start, "end_date": end})
    return float(t["Revenue"].iloc[0]) if not t.empty else 0.0


def _fetch(p: Period) -> dict:
    """Everything the report reads, fetched in parallel. The orders come one
    query per revenue-bridge period (7 days, 6 weeks or 6 months — which
    already hold this period and the one before it): shorter reads, less
    likely to be cancelled by the replica, each cached on its own (see
    _orders), so next week's report reuses five of its six weeks.

    The replica cancels long reads when it is busy ("conflict with recovery");
    run_query retries, and a read here gets one more round after a pause. If an
    older bridge period still fails, the bridge shows it as unavailable rather
    than the whole report failing; this period and the last must load."""
    import time
    from concurrent.futures import ThreadPoolExecutor
    from lib import db, feedback, footfall, queries

    spans = bridge_periods(p)
    timings: dict[str, float] = {}
    needed = {(p.start, p.end), (p.prev_start, p.prev_end)}

    def timed(name, fn, optional=False):
        def run():
            t = time.perf_counter()
            try:
                for attempt in (1, 2):
                    try:
                        return fn()
                    except Exception:
                        if attempt == 2:
                            if optional:
                                return None
                            raise
                        time.sleep(3)
            finally:
                timings[name] = round(time.perf_counter() - t, 2)
        return run

    def responses():
        try:
            return feedback.load()
        except feedback.FeedbackUnavailable:
            return None

    jobs = {
        **{f"orders {i}": (lambda s=s, e=e: _orders(s, e)) for i, (_, s, e) in enumerate(spans)},
        "corp_cur": lambda: _corporate(p.start, p.end),
        "corp_prev": lambda: _corporate(p.prev_start, p.prev_end),
        "corp_lines": lambda: db.run_query(queries.CORPORATE_LINES,
                                           {"start_date": p.start, "end_date": p.end}),
        "footfall": lambda: footfall.daily(p.start, p.end)[0],
        "responses": responses,
        "chats": lambda: db.run_query(queries.WA_INTERACTIONS,
                                      {"start_date": p.start, "end_date": p.end}),
    }
    # Cached calls made from these threads have no script context, so they
    # skip their spinners; the page shows one spinner over the whole build.
    # Eight at a time: the database pool holds ten connections.
    optional = {f"orders {i}" for i, (_, s, e) in enumerate(spans) if (s, e) not in needed}
    with ThreadPoolExecutor(max_workers=8) as pool:
        futures = {k: pool.submit(timed(k, f, optional=k in optional)) for k, f in jobs.items()}
        out = {k: f.result() for k, f in futures.items()}
    keys = [k for k in out if k.startswith("orders ")]
    parts = [out.pop(k) for k in keys]
    out["missing_spans"] = {spans[int(k.split()[1])][0] for k, x in zip(keys, parts) if x is None}
    parts = [x for x in parts if x is not None]
    out["orders"] = pd.concat([x for x in parts if not x.empty] or parts, ignore_index=True)
    out["spans"], out["timings"] = spans, timings
    return out


def build(p: Period) -> dict:
    """Every section of the report for period `p`, against the one before."""
    data = _fetch(p)
    orders = data["orders"].copy()
    orders["Date"] = pd.to_datetime(orders["Date"]).dt.date
    in_range = lambda s, e: orders[(orders["Date"] >= s) & (orders["Date"] <= e)]
    cur_raw, prev_raw = in_range(p.start, p.end), in_range(p.prev_start, p.prev_end)
    not_customers = placeholders(cur_raw, prev_raw)
    cur_df, prev_df = prepare(cur_raw, not_customers), prepare(prev_raw, not_customers)
    cur, prev = headline(cur_df, data["corp_cur"]), headline(prev_df, data["corp_prev"])

    out = {"period": p, "cur": cur, "prev": prev, "timings": data["timings"]}
    out["executive"], out["flags"], out["exec_note"] = executive(cur, prev, p)
    out["cust_metrics"], out["frequency"], out["channels"], out["cust_note"] =         customer_metrics(cur_df, prev_df, cur, prev)
    out["stores"], out["channel_mix"] = store_performance(cur_df, prev_df)
    out["retention"] = retention_table(cur_df, prev_df)
    (out["gender_ratio"], out["gender"], out["gender_shop"], out["gender_note"],
     out["gender_unknown"], gender_shift) = gender_split(cur_df, prev_df, p)

    foot = data["footfall"]
    out["traffic"], out["traffic_note"] = traffic(foot, cur_df) if not foot.empty else         (pd.DataFrame(), "No footfall was counted in this period.")

    # Revenue bridge: the last 7 days, 6 weeks or 6 months.
    window = prepare(orders, not_customers)
    rows, last = [], None
    for label, s, e in data["spans"]:
        if label in data["missing_spans"]:
            # The replica wouldn't serve this period even after retries.
            rows.append((label, np.nan, np.nan, np.nan, "unavailable"))
            continue
        part = window[(window["Date"] >= s) & (window["Date"] <= e)]
        rev = part["Revenue"].sum()
        custs = part.loc[part["Sale"] & part["Key"].notna(), "Key"].nunique()
        trend = "→" if last is None or abs(rev - last) < 1e-6 else ("▲" if rev > last else "▼")
        rows.append((label, round(rev, 2), custs, round(rev / custs, 2) if custs else 0.0, trend))
        last = rev
    out["bridge"] = pd.DataFrame(rows, columns=["Period", "Revenue", "Customers", "Avg Spend", "Trend"])
    out["revenue_kpis"] = pd.DataFrame([
        ("Revenue", money(cur["Revenue"]), change(cur["Revenue"], prev["Revenue"], money_fmt=True)),
        ("Avg Rev/Txn", money(cur["Avg Rev/Txn"]), change(cur["Avg Rev/Txn"], prev["Avg Rev/Txn"], money_fmt=True)),
        ("Avg Spend/Customer", money(cur["Avg Spend/Customer"]),
         change(cur["Avg Spend/Customer"], prev["Avg Spend/Customer"], money_fmt=True)),
        ("Avg Orders/Customer", f"{cur['Avg Orders']:.2f}", change(cur["Avg Orders"], prev["Avg Orders"])),
        ("Corporate Revenue", money(cur["Corporate"]), change(cur["Corporate"], prev["Corporate"], money_fmt=True)),
        ("Overall Revenue (incl. Corporate)", money(cur["Overall"]),
         change(cur["Overall"], prev["Overall"], money_fmt=True)),
    ], columns=["Label", "Value", "Change"])
    out["corporate"] = pd.DataFrame([
        ("Corporate Revenue", money(cur["Corporate"]), money(prev["Corporate"]),
         change(cur["Corporate"], prev["Corporate"], money_fmt=True)),
        ("Overall Revenue", money(cur["Overall"]), money(prev["Overall"]),
         change(cur["Overall"], prev["Overall"], money_fmt=True)),
    ], columns=["Metric", "Current", "Previous", "Change"])
    lines = data["corp_lines"]
    jobs = lines[lines["Channel"] == "Corporate Sales"] if not lines.empty else lines
    out["corporate_clients"] = (
        jobs.groupby("Customer", as_index=False)
        .agg(Orders=("Order", "nunique"), Bags=("Bags", "sum"), Revenue=("Revenue", "sum"),
             **{"First Order": ("Date", "min"), "Last Order": ("Date", "max")})
        .sort_values("Revenue", ascending=False)) if not jobs.empty else pd.DataFrame()

    (out["quality"], out["store_numbers"], out["staff_numbers"], out["quality_note"],
     data_notes) = data_quality(cur_df, prev_df, not_customers)

    responses, chats = data["responses"], data["chats"]
    links = chats[chats["Feedback Sent"]] if not chats.empty else None
    out["feedback"] = feedback_section(responses, links, p)
    out["notes"] = key_notes(p, cur, prev, out["stores"], out["traffic"], gender_shift,
                             out["feedback"], data_notes)
    out["generated"] = datetime.now()
    return out


# ---------------------------------------------------------------- export ---

def title_for(p: Period) -> str:
    return {DAILY: "Daily", WEEKLY: "Weekly", MONTHLY: "Monthly"}[p.kind] + " Sales Performance Report"


def summary_rows(r: dict) -> pd.DataFrame:
    cur, prev, p = r["cur"], r["prev"], r["period"]
    vs = f"vs {p.prev_label}"
    rows = [
        ("Revenue", money(cur["Revenue"]), f"{change(cur['Revenue'], prev['Revenue'])} {vs}",
         tone(cur["Revenue"], prev["Revenue"])),
        ("Corporate Revenue", money(cur["Corporate"]),
         f"{change(cur['Corporate'], prev['Corporate'], money_fmt=True)} {vs}",
         tone(cur["Corporate"], prev["Corporate"])),
        ("Overall Revenue (incl. Corporate)", money(cur["Overall"]),
         f"{change(cur['Overall'], prev['Overall'], money_fmt=True)} {vs}",
         tone(cur["Overall"], prev["Overall"])),
        ("Total Customers", f"{cur['Customers']:,}", f"{change(cur['Customers'], prev['Customers'])} {vs}",
         tone(cur["Customers"], prev["Customers"])),
        ("Avg Spend/Customer", money(cur["Avg Spend/Customer"]),
         f"{change(cur['Avg Spend/Customer'], prev['Avg Spend/Customer'])} {vs}",
         tone(cur["Avg Spend/Customer"], prev["Avg Spend/Customer"])),
        ("Repeat Rate", f"{cur['Repeat Rate']:.1f}%",
         f"{change(cur['Repeat Rate'], prev['Repeat Rate'], pp=True)} {vs}",
         tone(cur["Repeat Rate"], prev["Repeat Rate"])),
        ("Data Quality", f"{cur['Data Quality']:.1f}%",
         f"{change(cur['Data Quality'], prev['Data Quality'], pp=True)} {vs}",
         tone(cur["Data Quality"], prev["Data Quality"])),
    ]
    return pd.DataFrame(rows, columns=["Label", "Value", "Change", "Tone"])


def to_xlsx(r: dict) -> bytes:
    """The report as a workbook laid out like the team's: one sheet per
    section, titled and in the same order."""
    import io
    from openpyxl.styles import Font

    p = r["period"]
    buf = io.BytesIO()
    sheets: list[tuple[str, list]] = []

    def block(title=None, frame=None, text=None):
        return ("title", title) if title else ("table", frame) if frame is not None else ("text", text)

    fb = r["feedback"]
    sheets = [
        ("Summary", [block(title_for(p)), block(text=p.label), block(text=f"Compared to: {p.prev_label}"),
                     block(text=f"Generated: {r['generated']:%d %b %Y, %H:%M}"),
                     block(title="Key Performance Indicators"), block(frame=summary_rows(r)),
                     block(text=f"This report covers {p.days} trading day(s).")]),
        ("Executive Summary", [block(title="1. Executive Summary"), block(frame=r["executive"]),
                               block(title="Flags")] + [block(text=f) for f in r["flags"]]
         + [block(text=r["exec_note"])]),
        ("Customer Metrics", [block(title="2. Customer Metrics"), block(frame=r["cust_metrics"]),
                              block(title="2.1 Purchase Frequency Distribution"), block(frame=r["frequency"]),
                              block(title="2.2 Channel Overview"), block(frame=r["channels"]),
                              block(text=r["cust_note"])]),
        ("Store Performance", [block(title="3. Store Performance"), block(frame=r["stores"]),
                               block(title="3.1 Channel Mix by Location"), block(frame=r["channel_mix"])]
         + ([block(title="3.2 Monthly Customer Retention by Location"), block(frame=r["retention"])]
            if p.kind == MONTHLY else [])),
        ("Gender Performance", [block(title="4. Gender Performance Analysis"),
                                block(text=f"Ratio: {r['gender_ratio']}"),
                                block(title="4.1 Overall Gender Split"), block(frame=r["gender"]),
                                block(title="4.2 Gender by Location"), block(frame=r["gender_shop"]),
                                block(text=r["gender_note"])]),
        ("Traffic", [block(title="5. Foot & Online Traffic Analysis"), block(frame=r["traffic"]),
                     block(text=r["traffic_note"])]),
        ("Revenue Analysis", [block(title="6. Revenue Analysis"), block(frame=r["revenue_kpis"]),
                              block(title="6.1 Revenue Bridge"), block(frame=r["bridge"]),
                              block(title="6.2 Corporate Revenue"), block(frame=r["corporate"]),
                              block(title="6.3 Top Corporate Clients"),
                              block(frame=r["corporate_clients"]) if not r["corporate_clients"].empty
                              else block(text="No data for this period.")]),
        ("Data Quality", [block(title="7. Data Quality (KYC)"), block(frame=r["quality"]),
                          block(title="7.1 Store & Attendant Numbers Used by Location"),
                          block(frame=r["store_numbers"]),
                          block(title="7.2 Store & Attendant Numbers on Customer Receipts"),
                          block(frame=r["staff_numbers"]),
                          block(text=r["quality_note"])]),
        ("Feedback", [block(title="8. Customer Feedback")]
         + ([block(frame=fb["summary"]), block(title="8.1 Store Breakdown"), block(frame=fb["by_shop"])]
            + ([block(title="8.2 Feedback Link Targets by Location"), block(frame=fb["targets"])]
               if p.kind != DAILY and fb["targets"] is not None else [])
            if fb else [block(text="The feedback sheet could not be read.")])),
        ("Summary & Key Notes", [block(title="9. Summary & Key Notes")]
         + [b for team, items in r["notes"].items()
            for b in [block(title=team),
                      block(text=f"{len(items)} item(s) flagged for {team} this period." if items
                            else f"No {team.lower()} flags this period.")]
            + [block(text=i) for i in items]]),
    ]
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, blocks in sheets:
            row = 0
            pd.DataFrame().to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            for kind, value in blocks:
                if kind == "table":
                    if value is None or value.empty:
                        ws.cell(row=row + 1, column=1, value="No data for this period.")
                        row += 2
                        continue
                    value.to_excel(writer, sheet_name=name, index=False, startrow=row)
                    for c in range(1, len(value.columns) + 1):
                        ws.cell(row=row + 1, column=c).font = Font(bold=True)
                    row += len(value) + 2
                else:
                    cell = ws.cell(row=row + 1, column=1, value=value)
                    if kind == "title":
                        cell.font = Font(bold=True, size=13)
                    row += 2 if kind == "title" else 1
            for column in ws.columns:
                width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
                ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 48)
    return buf.getvalue()
